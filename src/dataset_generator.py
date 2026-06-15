import os
import csv
import random
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document

# Define categories
CATEGORIES = [
    "Resume",
    "Invoice",
    "Medical Report",
    "Legal Document",
    "News Article",
    "Research Paper"
]

# Vocabulary and templates for text generation
NAMES = ["John Smith", "Jane Doe", "Alex Rivera", "Emily Chen", "Michael Scott", "Sarah Jenkins", "David Patel", "Rachel Green", "James Watson", "Sophia Al-Jamil"]
ROLES = ["Senior Software Engineer", "Data Scientist", "Product Manager", "DevOps Engineer", "Machine Learning Specialist", "UX Designer", "Marketing Manager"]
TECH_SKILLS = ["Python", "SQL", "Scikit-Learn", "TensorFlow", "PyTorch", "Docker", "AWS", "Git", "Kubernetes", "React", "Node.js", "Java", "C++", "FastAPI"]
UNIVERSITIES = ["MIT", "Stanford University", "UC Berkeley", "Carnegie Mellon", "Harvard University", "Oxford University", "University of Toronto"]
COMPANIES = ["TechCorp Industries", "Innovate LLC", "Stark Enterprises", "Wayne Tech", "Globex Corporation", "Apex Systems", "Acme Inc"]

VENDORS = ["Global Solutions Inc.", "Initech Corp", "Apex Consulting", "Cyberdyne Systems", "Tyrell Biotech", "Umbrella Corp"]
CLIENTS = ["Hooli Inc.", "Dunder Mifflin", "Vandelay Industries", "Soylent Corp", "Pied Piper", "Bluth Company"]
ITEMS = [
    ("Software Development Services", 80, 100),
    ("Cloud Architecture Consulting", 40, 150),
    ("Machine Learning Engineering", 60, 120),
    ("UI/UX Prototype Design", 30, 90),
    ("Database Migration & Optimization", 50, 110),
    ("Security Audit and Penetration Testing", 25, 200)
]

SYMPTOMS = [
    "persistent non-productive cough, fatigue, and low-grade fever for 4 days",
    "acute epigastric abdominal pain, nausea, and occasional vomiting",
    "recurrent throbbing headache in the temporal region with photophobia",
    "chronic lower back pain radiating down the left leg, aggravated by sitting",
    "shortness of breath on exertion, mild bilateral ankle edema, and palpitations",
    "sore throat, difficulty swallowing, nasal congestion, and body aches"
]
DIAGNOSES = [
    "Mild Acute Bronchitis - secondary to viral infection",
    "Gastroesophageal Reflux Disease (GERD) with mild gastritis",
    "Migraine headache without aura - stress-induced",
    "Lumbar radiculopathy - suspected L4-L5 disc bulge",
    "Hypertension Stage 1 and chronic venous insufficiency",
    "Acute Pharyngitis - symptomatic treatment initiated"
]
MEDICATIONS = [
    "Amoxicillin 500mg - three times daily for 7 days",
    "Omeprazole 20mg - daily before breakfast for 14 days",
    "Sumatriptan 50mg - as needed at onset of headache",
    "Ibuprofen 400mg - every 6 hours as needed for pain",
    "Lisinopril 10mg - daily in the morning",
    "Acetaminophen 500mg - every 4-6 hours for fever/aches"
]

LEGAL_TYPES = ["MUTUAL NONDISCLOSURE AGREEMENT", "COMMERCIAL LEASE AGREEMENT", "INDEPENDENT CONTRACTOR AGREEMENT", "SOFTWARE LICENSE AGREEMENT"]
STATES = ["State of New York", "State of California", "State of Delaware", "State of Texas", "Commonwealth of Massachusetts"]
CLAUSES = [
    "The Recipient shall keep the Confidential Information strictly confidential and shall not disclose it to any third party without prior written consent.",
    "This Agreement shall be governed by and construed in accordance with the laws of the jurisdiction specified, excluding its conflict of law principles.",
    "Either party may terminate this Agreement upon thirty (30) days written notice to the other party in the event of a material breach of terms.",
    "The Contractor is an independent contractor, and nothing in this Agreement shall create an employer-employee relationship, partnership, or joint venture.",
    "If any provision of this Agreement is held to be invalid or unenforceable, the remaining provisions shall continue in full force and effect."
]

NEWS_HEADLINES = [
    "City Council Approves Multibillion Green Transit Initiative",
    "Tech Startup Unveils Next-Generation Solid-State Battery",
    "Local Football Club Clinches Championship in Dramatic Shootout",
    "Central Bank Raises Benchmark Interest Rate to Curb Inflation",
    "Global Summit Addresses Cybersecurity Threats in Critical Infrastructure",
    "Leading Research Institute Announces Breakthrough in Quantum Computing"
]
CITIES = ["New York", "London", "Tokyo", "Berlin", "San Francisco", "Sydney", "Singapore"]
NEWS_QUOTES = [
    "We believe this policy is the first step toward a more sustainable and economically vibrant future.",
    "Our team has worked tirelessly for three years to bring this revolutionary technology to the commercial market.",
    "The dedication shown by these players represents the highest standard of sportsmanship and team effort.",
    "We must act decisively to maintain price stability while supporting sustainable economic growth.",
    "No single nation can defend against these distributed threats alone; collaboration is absolutely crucial."
]

RESEARCH_TITLES = [
    "An Empirical Evaluation of Document Classification Using Machine Learning",
    "Optimizing Transformer Architectures for Resource-Constrained Environments",
    "A Hybrid Approach to Natural Language Processing using TF-IDF and Neural Networks",
    "Analyzing the Impact of Hyperparameter Tuning on SVM and Logistic Regression Classifiers",
    "Robust Feature Selection Methodologies for Text Mining and Classification Tasks",
    "Evaluating Classical Machine Learning Algorithms for Document Genre Sorting"
]

def generate_synthetic_text(category):
    """Generates a text string typical of the requested category."""
    if category == "Resume":
        name = random.choice(NAMES)
        role = random.choice(ROLES)
        skills = random.sample(TECH_SKILLS, k=5)
        uni = random.choice(UNIVERSITIES)
        degree = random.choice(["B.S. in Computer Science", "M.S. in Data Science", "Ph.D. in Artificial Intelligence", "B.A. in Software Engineering"])
        company1, company2 = random.sample(COMPANIES, k=2)
        
        text = f"""
        {name}
        Email: {name.lower().replace(' ', '.')}@example.com | Phone: (555) 019-2834 | Address: San Francisco, CA
        
        PROFESSIONAL SUMMARY
        Highly motivated and results-driven {role} with over 5 years of experience in designing, building, and deploying scalable software systems. Expertise in machine learning pipelines, full-stack architecture, and agile methodologies. Proven record of optimization and leadership.
        
        AREAS OF EXPERTISE
        Technical Skills: {', '.join(skills)}, SQL databases, RESTful APIs, Git version control, software design patterns, cloud computing.
        Soft Skills: Project management, technical mentoring, collaborative team building, agile scrum, system design thinking.
        
        PROFESSIONAL EXPERIENCE
        Senior Analyst / Developer | {company1} | 2022 - Present
        - Engineered high-throughput data pipelines using Python, improving processing efficiency by 35%.
        - Designed and deployed custom machine learning classifiers, achieving 92% classification accuracy on unstructured text datasets.
        - Led a team of 4 engineers in migrating legacy infrastructure to containerized microservices in AWS.
        
        Software Engineer | {company2} | 2019 - 2022
        - Collaborated with product owners to define application specifications and implement core APIs.
        - Maintained 99.9% uptime for business-critical payment platforms by redesigning caching layers.
        - Wrote comprehensive unit tests and automated CI/CD integration scripts.
        
        EDUCATION
        {degree}
        {uni} | Graduation: June 2019
        Relevant Coursework: Data Structures, Machine Learning, Database Systems, Natural Language Processing, Software Architecture.
        """
    
    elif category == "Invoice":
        vendor = random.choice(VENDORS)
        client = random.choice(CLIENTS)
        inv_num = f"INV-2026-{random.randint(1000, 9999)}"
        item1, item2 = random.sample(ITEMS, k=2)
        
        qty1, rate1 = item1[1], item1[2]
        qty2, rate2 = item2[1], item2[2]
        subtotal = (qty1 * rate1) + (qty2 * rate2)
        tax = round(subtotal * 0.0825, 2)
        total = subtotal + tax
        
        text = f"""
        INVOICE
        {vendor}
        100 Corporate Parkway, Suite 400
        Email: billing@{vendor.lower().replace(' ', '').replace('.', '')}.com
        
        INVOICE TO:
        {client}
        Attn: Accounts Payable
        456 Business Road, Tech Center
        
        INVOICE DETAILS:
        Invoice Number: {inv_num}
        Invoice Date: June 15, 2026
        Due Date: July 15, 2026
        Payment Terms: Net 30
        
        LINE ITEMS:
        1. Description: {item1[0]}
           Quantity: {qty1} hours | Unit Price: ${rate1:.2f} | Total: ${qty1*rate1:.2f}
        2. Description: {item2[0]}
           Quantity: {qty2} hours | Unit Price: ${rate2:.2f} | Total: ${qty2*rate2:.2f}
           
        SUMMARY:
        Subtotal: ${subtotal:.2f}
        Tax Rate: 8.25%
        Estimated Tax: ${tax:.2f}
        TOTAL AMOUNT DUE: ${total:.2f}
        
        Payment Instructions:
        Please transfer the total amount to Bank of America, Account: 9876-5432-1098, Routing: 021000021.
        For billing questions, contact support@{vendor.lower().replace(' ', '').replace('.', '')}.com. Thank you for your business!
        """
        
    elif category == "Medical Report":
        patient = random.choice(NAMES)
        dob = f"{random.randint(1, 12)}/{random.randint(1, 28)}/{random.randint(1960, 2005)}"
        symptom = random.choice(SYMPTOMS)
        diagnosis = random.choice(DIAGNOSES)
        medication = random.choice(MEDICATIONS)
        systolic = random.randint(110, 145)
        diastolic = random.randint(70, 95)
        pulse = random.randint(60, 95)
        temp = round(random.uniform(97.8, 101.5), 1)
        
        text = f"""
        CLINICAL ENCOUNTER SUMMARY & MEDICAL REPORT
        St. Jude Memorial Hospital | Outpatient Clinic
        
        PATIENT DEMOGRAPHICS:
        Patient Name: {patient}
        Date of Birth: {dob}
        Gender: {random.choice(['Male', 'Female', 'Other'])}
        Date of Encounter: June 15, 2026
        Primary Physician: Dr. Elizabeth Blackwell, MD
        
        CLINICAL PRESENTATION & CHIEF COMPLAINT:
        The patient presents with {symptom}. Symptoms have been persistent, affecting daily routines. No history of travel or exposure to known sick contacts.
        
        VITAL SIGNS:
        - Blood Pressure: {systolic}/{diastolic} mmHg (sitting)
        - Heart Rate: {pulse} bpm (regular)
        - Temperature: {temp} F (oral)
        - Respiratory Rate: {random.randint(14, 20)} breaths/min
        - SpO2: {random.randint(96, 99)}% on room air
        
        PHYSICAL EXAMINATION:
        General: Well-developed, well-nourished individual in mild distress.
        Chest/Lungs: Clear to auscultation bilaterally, no wheezes or rales.
        Cardiovascular: Normal S1, S2, regular rhythm, no murmurs detected.
        Abdomen: Soft, non-distended, mild tenderness as noted in report.
        
        ASSESSMENT / DIAGNOSIS:
        {diagnosis}.
        
        TREATMENT PLAN & PRESCRIPTION:
        1. Rx: {medication}. Take exactly as directed.
        2. Rest and adequate hydration. Avoid strenuous physical activity for 72 hours.
        3. Follow-up in clinic in 10 days if symptoms do not improve.
        4. Patient instructed to present to the Emergency Department immediately if chest pain, severe dyspnea, or high fever develops.
        
        Dr. Elizabeth Blackwell, MD
        License Number: MD-9827364-CA
        """
        
    elif category == "Legal Document":
        legal_type = random.choice(LEGAL_TYPES)
        state = random.choice(STATES)
        company1, company2 = random.sample(COMPANIES, k=2)
        clause1, clause2 = random.sample(CLAUSES, k=2)
        
        text = f"""
        {legal_type}
        
        THIS AGREEMENT is entered into this 15th day of June, 2026 (the "Effective Date"), by and between:
        - {company1}, a corporation organized and existing under the laws of the {state}, with its principal place of business at 123 Corporate Blvd (hereinafter "Disclosing Party"), and
        - {company2}, a limited liability company organized and existing under the laws of the {state}, with its principal place of business at 789 Partner St (hereinafter "Receiving Party").
        
        WHEREAS, the parties desire to enter into discussions regarding a potential business transaction, collaboration, or relationship (the "Purpose"); and
        
        WHEREAS, in connection with the Purpose, the Disclosing Party may disclose certain proprietary, confidential, or trade secret information to the Receiving Party;
        
        NOW, THEREFORE, in consideration of the mutual covenants contained herein and other good and valuable consideration, the parties agree as follows:
        
        1. Confidentiality obligations: {clause1}
        2. Governing law: {clause2}
        3. Term: This Agreement shall commence on the Effective Date and remain in effect for a period of two (2) years, unless terminated earlier by mutual written agreement.
        4. Remedies: The Receiving Party acknowledges that a breach of this Agreement may cause irreparable harm for which monetary damages would be inadequate, and agrees that the Disclosing Party shall be entitled to seek injunctive relief.
        
        IN WITNESS WHEREOF, the parties hereto have executed this Agreement by their authorized representatives as of the Effective Date written above.
        
        For Disclosing Party ({company1}):
        By: ___________________________
        Name: John Smith, CEO
        
        For Receiving Party ({company2}):
        By: ___________________________
        Name: Jane Doe, Managing Director
        """
        
    elif category == "News Article":
        headline = random.choice(NEWS_HEADLINES)
        city = random.choice(CITIES)
        quote = random.choice(NEWS_QUOTES)
        
        text = f"""
        THE DAILY GAZETTE
        Edition: National News | Date: June 15, 2026
        
        {headline.upper()}
        
        BYLINE: Associated Press
        DATELINE: {city} — June 15, 2026
        
        Officials confirmed yesterday that new administrative measures have been approved to address ongoing infrastructure challenges. The decision followed a heated debate in the assembly, culminating in a bipartisan consensus that seeks to stimulate local economic growth and foster sustainability.
        
        Local representatives expressed optimism about the outcomes. "This is a historic milestone for our community," said a government spokesperson, noting that the long-term impact on employment and transport efficiency would be felt almost immediately.
        
        Critics, however, raise concerns about the short-term financial burdens and potential delays in project execution. Industry analysts suggest that close oversight is required to prevent cost overruns, which have plagued similar initiatives in neighboring districts.
        
        In a public press conference, the lead director stated: "{quote}"
        
        Further updates are expected as planning committees hold public hearings next month to finalize timelines. Local businesses are encouraged to submit applications for public contracts starting next week.
        
        © 2026 The Daily Gazette Media Group. All rights reserved.
        """
        
    elif category == "Research Paper":
        title = random.choice(RESEARCH_TITLES)
        uni1, uni2 = random.sample(UNIVERSITIES, k=2)
        
        text = f"""
        {title}
        
        Author Names: Dr. Arthur Pendelton, Prof. Beverly Crusher
        Department of Computer Science and Intelligent Systems
        {uni1} & {uni2}
        
        ABSTRACT
        In this paper, we explore the performance of classical machine learning models for the task of document classification. Using TF-IDF (Term Frequency-Inverse Document Frequency) vectorization, we extract key lexical features from structured and unstructured documents. We evaluate three popular classification algorithms: Multinomial Naive Bayes, Logistic Regression, and Support Vector Machines (SVM). Experimental results show that while SVM achieves the highest overall accuracy (97.4%), Multinomial Naive Bayes offers superior training speed and handles sparse high-dimensional data exceptionally well.
        
        1. INTRODUCTION
        Document classification is a fundamental problem in Natural Language Processing (NLP) with applications in legal document analysis, medical record sorting, invoice routing, and automated news categorization. The volume of digital text requires automated techniques to filter and route records.
        
        2. METHODOLOGY AND DATA PREPROCESSING
        We compile a corpus of document classes: Resumes, Invoices, Medical Reports, Legal Documents, News Articles, and Research Papers.
        The text preprocessing pipeline involves:
        - Case normalization (converting all characters to lowercase).
        - Stripping special characters and digits.
        - Vectorization via TF-IDF:
          TF-IDF(t, d, D) = TF(t, d) * IDF(t, D)
        
        3. EXPERIMENTAL SETUP AND RESULTS
        Our model was evaluated using a stratified 80/20 train-test split. Table 1 outlines the performance metrics across all models.
        - Support Vector Machine: Accuracy 97.4%, Precision 97.1%, Recall 97.4%
        - Logistic Regression: Accuracy 96.2%, Precision 96.3%, Recall 96.2%
        - Multinomial Naive Bayes: Accuracy 94.8%, Precision 95.0%, Recall 94.8%
        
        4. CONCLUSION & FUTURE WORK
        We demonstrated that classical machine learning algorithms paired with TF-IDF features remain highly competitive baselines for domain-specific text classification. Future investigations will evaluate dense embeddings.
        
        REFERENCES
        [1] Smith, J. and Jones, A. "Text Classification in the Digital Era." Journal of NLP Research, 2024.
        [2] Bengio, Y. "Representations for Natural Language." IEEE Transactions, 2021.
        [3] Cover, T. and Thomas, J. "Elements of Information Retrieval." 2018.
        """
        
    else:
        text = "Sample document content."
        
    return text.strip()

def generate_dataset_csv(output_csv_path, num_samples_per_class=60):
    """
    Generates a large synthetic dataset and saves it to a CSV file.
    """
    os.makedirs(os.path.dirname(output_csv_path), exist_ok=True)
    
    with open(output_csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["text", "category"])
        
        for category in CATEGORIES:
            # Generate multiple unique instances
            for _ in range(num_samples_per_class):
                text_sample = generate_synthetic_text(category)
                writer.writerow([text_sample, category])
                
    print(f"Generated synthetic dataset with {num_samples_per_class * len(CATEGORIES)} samples at {output_csv_path}")

def create_sample_pdf(text, output_path):
    """Creates a beautifully styled PDF from text using reportlab."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = SimpleDocTemplate(output_path, pagesize=letter,
                            rightMargin=54, leftMargin=54,
                            topMargin=54, bottomMargin=54)
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1A365D"),
        spaceAfter=15,
        alignment=1 # Center
    )
    
    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#2B6CB0"),
        spaceBefore=10,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#2D3748"),
        spaceAfter=8
    )
    
    story = []
    
    lines = text.split("\n")
    is_first = True
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue
            
        # Very naive parser to make the PDF look formatted
        if is_first:
            story.append(Paragraph(stripped, title_style))
            story.append(Spacer(1, 10))
            is_first = False
        elif stripped.isupper() and len(stripped) < 50:
            story.append(Paragraph(stripped, heading_style))
        elif stripped.startswith("- "):
            story.append(Paragraph(f"&bull; {stripped[2:]}", body_style))
        else:
            story.append(Paragraph(stripped, body_style))
            
    doc.build(story)

def create_sample_docx(text, output_path):
    """Creates a DOCX from text using python-docx."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()
    
    lines = text.split("\n")
    is_first = True
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        if is_first:
            doc.add_heading(stripped, level=0)
            is_first = False
        elif stripped.isupper() and len(stripped) < 50:
            doc.add_heading(stripped, level=1)
        else:
            doc.add_paragraph(stripped)
            
    doc.save(output_path)

def create_sample_txt(text, output_path):
    """Creates a plain text file."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)

def generate_all_samples(samples_dir):
    """Generates 6 mock documents, one for each category, in various formats."""
    os.makedirs(samples_dir, exist_ok=True)
    
    # We will generate one PDF, DOCX, and TXT for alternating categories
    formats = ["pdf", "docx", "txt", "pdf", "docx", "txt"]
    
    for category, file_format in zip(CATEGORIES, formats):
        file_name = f"sample_{category.lower().replace(' ', '_')}.{file_format}"
        file_path = os.path.join(samples_dir, file_name)
        
        # Generate raw text for the file
        raw_text = generate_synthetic_text(category)
        
        if file_format == "pdf":
            create_sample_pdf(raw_text, file_path)
        elif file_format == "docx":
            create_sample_docx(raw_text, file_path)
        elif file_format == "txt":
            create_sample_txt(raw_text, file_path)
            
        print(f"Generated sample file: {file_path}")

if __name__ == "__main__":
    # Test generation if run directly
    generate_dataset_csv("data/document_dataset.csv", num_samples_per_class=10)
    generate_all_samples("data/samples")
