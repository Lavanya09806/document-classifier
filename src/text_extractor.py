import re
import os
from pypdf import PdfReader
from docx import Document

def clean_text(text):
    """
    Cleans raw text: removes multiple spaces, replaces newlines with spaces,
    and performs basic normalization.
    """
    if not text:
        return ""
    # Replace newlines and carriage returns with space
    text = text.replace("\r", " ").replace("\n", " ")
    # Keep alphanumeric characters, spaces, and basic punctuation
    text = re.sub(r"[^\w\s\.,;:!?\-\'\"@#\$%&\*\(\)]", "", text)
    # Remove multiple spaces
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def extract_text_from_txt(file_path):
    """Extracts text from a plain text file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        raise ValueError(f"Error reading TXT file: {str(e)}")

def extract_text_from_docx(file_path):
    """Extracts text from a Word document (.docx)."""
    try:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        # Include table cells if any
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_text.append(cell.text)
        
        full_text = "\n".join(paragraphs + table_text)
        return full_text
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {str(e)}")

def extract_text_from_pdf(file_path):
    """Extracts text from a PDF document."""
    try:
        reader = PdfReader(file_path)
        text_list = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_list.append(page_text)
        return "\n".join(text_list)
    except Exception as e:
        raise ValueError(f"Error reading PDF file: {str(e)}")

def extract_text(file_path):
    """
    Determines file type by extension and extracts text.
    Returns cleaned text.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    _, ext = os.path.splitext(file_path.lower())
    
    if ext == ".txt":
        raw_text = extract_text_from_txt(file_path)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(file_path)
    elif ext == ".pdf":
        raw_text = extract_text_from_pdf(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")
        
    return clean_text(raw_text)
