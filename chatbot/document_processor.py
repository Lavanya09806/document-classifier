import os
from typing import List, Dict, Any
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
import pypdf
import docx

def extract_text_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """Extracts text from a PDF file page by page with metadata."""
    pages = []
    with open(file_path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages.append({
                    "text": text.strip(),
                    "metadata": {
                        "source": os.path.basename(file_path),
                        "page": page_num + 1,
                        "type": "pdf"
                    }
                })
    return pages

def extract_text_from_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extracts text from a DOCX file, including tables."""
    doc = docx.Document(file_path)
    full_text = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    
    # Extract table contents nicely
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
            
    text = "\n".join(full_text)
    if text.strip():
        return [{
            "text": text.strip(),
            "metadata": {
                "source": os.path.basename(file_path),
                "type": "docx"
            }
        }]
    return []

def extract_text_from_txt(file_path: str) -> List[Dict[str, Any]]:
    """Extracts text from a plain text or markdown file."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    if text.strip():
        return [{
            "text": text.strip(),
            "metadata": {
                "source": os.path.basename(file_path),
                "type": "txt"
            }
        }]
    return []

def load_document(file_path: str) -> List[Dict[str, Any]]:
    """Determines file type and extracts its content with metadata."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md", ".csv", ".json", ".xml", ".py"]:
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

def process_documents(file_paths: List[str], chunk_size: int = 1000, chunk_overlap: int = 200) -> List[Document]:
    """
    Loads multiple files, extracts text, chunks them using RecursiveCharacterTextSplitter,
    and returns a list of LangChain Document objects.
    """
    all_pages = []
    for path in file_paths:
        try:
            pages = load_document(path)
            all_pages.extend(pages)
        except Exception as e:
            print(f"Error loading document {path}: {str(e)}")
            
    # Text splitter setting
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len
    )
    
    documents = []
    for page in all_pages:
        chunks = text_splitter.split_text(page["text"])
        for i, chunk in enumerate(chunks):
            metadata = page["metadata"].copy()
            metadata["chunk"] = i
            # Also keep a clean reference display string
            source_display = metadata["source"]
            if "page" in metadata:
                source_display += f" (Page {metadata['page']})"
            metadata["source_display"] = source_display
            
            documents.append(Document(page_content=chunk, metadata=metadata))
            
    return documents
